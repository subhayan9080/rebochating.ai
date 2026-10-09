"""File parsing + website crawling for knowledge sources."""
from __future__ import annotations

import os
import re
import urllib.parse
from collections import deque

MAX_UPLOAD_BYTES = 8 * 1024 * 1024
MAX_CRAWL_PAGES = 25
FETCH_TIMEOUT = 15
UA = {"User-Agent": "RelayBot/1.0 (+https://relay.local; knowledge indexer)"}

_WS = re.compile(r"\s+")


def clean_text(text: str) -> str:
    text = text.replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def parse_text_file(path: str) -> str:
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            with open(path, "r", encoding=enc) as f:
                return clean_text(f.read())
        except (UnicodeDecodeError, ValueError):
            continue
    raise ValueError("Could not decode text file")


def parse_pdf(path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return clean_text("\n\n".join(pages))


def parse_docx(path: str) -> str:
    import docx

    doc = docx.Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return clean_text("\n".join(parts))


def parse_csv_like(path: str) -> str:
    import csv

    try:
        with open(path, "r", encoding="utf-8", newline="") as f:
            sample = f.read(4096)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=[",", ";", "\t", "|"])
            except Exception:
                dialect = csv.excel
            rows = [" ".join(cell.strip() for cell in row if cell.strip())
                    for row in csv.reader(f, dialect)]
    except (UnicodeDecodeError, ValueError):
        rows = [parse_text_file(path)]
    return clean_text("\n".join(r for r in rows if r))


def parse_uploaded_file(path: str, filename: str) -> str:
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".pdf":
        return parse_pdf(path)
    if ext in (".docx",):
        return parse_docx(path)
    if ext in (".csv", ".tsv"):
        return parse_csv_like(path)
    if ext in (".txt", ".md", ".markdown", ".json", ".html", ".htm"):
        return parse_text_file(path)
    raise ValueError(f"Unsupported file type: {ext or '(none)'} — use PDF, DOCX, TXT, MD, CSV or JSON")


def same_site(a: str, b: str) -> bool:
    try:
        pa, pb = urllib.parse.urlparse(a), urllib.parse.urlparse(b)
        host = lambda h: h.lower().lstrip("www.")  # noqa: E731
        return host(pa.netloc) == host(pb.netloc) and pa.scheme in ("http", "https")
    except Exception:
        return False


def fetch_page_text(url: str) -> tuple[str, list[str]]:
    """Return (visible text, internal links)."""
    from bs4 import BeautifulSoup

    import requests

    resp = requests.get(url, headers=UA, timeout=FETCH_TIMEOUT)
    resp.raise_for_status()
    ctype = resp.headers.get("Content-Type", "")
    if "html" not in ctype and "text" not in ctype and not ctype.startswith("text/"):
        return "", []
    soup = BeautifulSoup(resp.text, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "nav", "footer", "header", "form", "aside"]):
        tag.decompose()
    text = clean_text(soup.get_text(separator="\n"))
    links: list[str] = []
    for a in soup.find_all("a", href=True):
        href = urllib.parse.urljoin(url, a["href"].strip())
        href = href.split("#")[0].rstrip("/")
        if href.startswith(("http://", "https://")) and same_site(url, href):
            if not re.search(r"\.(pdf|zip|png|jpe?g|gif|mp4|css|js|ico|svg|woff2?)($|\?)", href):
                links.append(href)
    return text, list(dict.fromkeys(links))


def crawl_website(start_url: str, max_pages: int = 8) -> list[dict]:
    """Breadth-first crawl, same site only. Returns [{url, title, text}]."""
    from bs4 import BeautifulSoup

    import requests

    if not re.match(r"^https?://", start_url, re.I):
        start_url = "https://" + start_url
    max_pages = max(1, min(int(max_pages or 8), MAX_CRAWL_PAGES))

    # Try sitemap.xml first for a smarter page list
    seeds: list[str] = [start_url.rstrip("/")]
    try:
        base = f"{urllib.parse.urlparse(start_url).scheme}://{urllib.parse.urlparse(start_url).netloc}"
        sm = requests.get(base + "/sitemap.xml", headers=UA, timeout=FETCH_TIMEOUT)
        if sm.ok and "<url" in sm.text[:2000]:
            soup = BeautifulSoup(sm.text, "xml")
            for loc in soup.find_all("loc")[: max_pages * 3]:
                u = (loc.get_text() or "").strip().split("#")[0].rstrip("/")
                if u.startswith(("http://", "https://")) and same_site(start_url, u):
                    seeds.append(u)
            seeds = list(dict.fromkeys(seeds))[:max_pages]
    except Exception:
        pass

    pages: list[dict] = []
    seen = set()
    queue: deque[str] = deque(seeds)
    while queue and len(pages) < max_pages:
        url = queue.popleft()
        if url in seen:
            continue
        seen.add(url)
        try:
            text, links = fetch_page_text(url)
        except Exception:
            continue
        if len(text) < 200:
            continue
        pages.append({"url": url, "title": url.split("//", 1)[-1][:80], "text": text})
        for link in links:
            if link not in seen and len(pages) + len(queue) < max_pages + 5:
                queue.append(link)
    if not pages:
        raise ValueError("No readable pages found. Check the URL is public and not JS-only.")
    return pages
